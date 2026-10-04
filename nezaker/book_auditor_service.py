import os
import re
import json
import hashlib
import time
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

import database
import gemini_service
import book_snippet_generator

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
AUDIT_DOCS_DIR = os.path.join(BASE_DIR, "reference_books")
os.makedirs(AUDIT_DOCS_DIR, exist_ok=True)

def normalize_text_hash(text: str) -> str:
    cleaned = re.sub(r'[^a-zA-Z0-9\u0600-\u06FF]', '', (text or '').lower())
    return hashlib.sha256(cleaned.encode('utf-8')).hexdigest()

def verify_single_question_against_book(
    book_id: int,
    question_text: str,
    options_text: str = "",
    proposed_answer: str = "",
    conversation_history: list = None,
    force_web_search: bool = False
) -> dict:
    """
    Verifies a question strictly against the book.
    If cached and confirmed convinced, returns cached.
    Otherwise searches candidate pages and invokes Gemini.
    """
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM reference_books WHERE id=?", (book_id,))
    book = cur.fetchone()
    if not book:
        conn.close()
        return {"status": "error", "message": "الكتاب المرجعي غير موجود"}

    q_hash = normalize_text_hash(question_text)

    # 1. Check cache if not forcing web search and not in consultation mode
    if not force_web_search and not conversation_history:
        cur.execute(
            "SELECT * FROM book_verified_qa WHERE book_id=? AND question_hash=? AND is_convinced=1",
            (book_id, q_hash)
        )
        cached = cur.fetchone()
        if cached:
            conn.close()
            return {
                "status": "cached",
                "verified_answer": cached['answer_text'],
                "explanation": cached['explanation'],
                "page_number": cached['page_number'],
                "snippet_image_url": cached['snippet_image_path'],
                "is_convinced": 1,
                "cached": True
            }

    pdf_path = book['file_path']
    book_title = book['title']

    # 2. Search relevant pages in the book
    candidate_pages = book_snippet_generator.search_relevant_pages(book_id, question_text + " " + options_text, max_pages=6)

    # If user explicitly requested web search
    if force_web_search:
        prompt = f"""
        You are an elite medical reference validator.
        The student is asking a medical question that was not found in their primary textbook "{book_title}".
        Search established medical and scientific consensus (PubMed, Medscape, UpToDate, authoritative medical literature).

        QUESTION:
        {question_text}

        OPTIONS (if any):
        {options_text}

        PROPOSED ANSWER (if any):
        {proposed_answer}

        CRITICAL INSTRUCTIONS:
        1. State whether the proposed answer is correct, or what the correct answer is according to authoritative medical consensus.
        2. Provide a thorough, crystal-clear explanation in Egyptian medical student register (Arabic with proper Medical English terms for all anatomical/clinical entities).
        3. You MUST provide the specific scientific reference/source (e.g. journal name, medical guideline, standard textbook, or URL if known).

        Respond in strict JSON format:
        {{
            "answer_status": "web_verified",
            "verified_answer": "...",
            "is_proposed_correct": true/false,
            "explanation_ar": "...",
            "source_reference": "...",
            "evidence_quote": "..."
        }}
        """
        client = gemini_service.get_gemini_client()
        for model in gemini_service.MODEL_PREFERENCES:
            try:
                response = client.models.generate_content(
                    model=model,
                    contents=prompt,
                    config=gemini_service.types.GenerateContentConfig(
                        temperature=0.2,
                        response_mime_type="application/json"
                    )
                )
                if response and response.text:
                    if hasattr(response, 'usage_metadata') and response.usage_metadata:
                        gemini_service.log_token_usage(
                            model,
                            response.usage_metadata.prompt_token_count or 0,
                            response.usage_metadata.candidates_token_count or 0,
                            response.usage_metadata.total_token_count or 0,
                            "book_auditor_web_search"
                        )
                    data = json.loads(gemini_service.clean_json_text(response.text))
                    conn.close()
                    return {
                        "status": "web_verified",
                        "verified_answer": data.get("verified_answer", proposed_answer),
                        "is_proposed_correct": data.get("is_proposed_correct", True),
                        "explanation": data.get("explanation_ar", ""),
                        "source_reference": data.get("source_reference", "المصادر الطبية المعتمدة"),
                        "page_number": 0,
                        "snippet_image_url": "",
                        "cached": False
                    }
            except Exception as e:
                continue

        conn.close()
        return {"status": "error", "message": "تعذر البحث عبر الإنترنت حالياً"}

    # 3. Regular Book Verification: Strictly based on book pages
    if not candidate_pages:
        conn.close()
        return {
            "status": "not_found",
            "message": "لم أتمكن من العثور على أي صفحات تتناول هذا السؤال في الكتاب المرفوع.",
            "can_search_web": True
        }

    pages_context = ""
    for score, p_num, p_text in candidate_pages:
        pages_context += f"\n--- PAGE {p_num} ---\n{p_text[:2500]}\n"

    prompt = f"""
    You are a strict Medical Textbook Reference Auditor.
    Your mission is to verify the student's question SOLELY based on the provided textbook excerpts from "{book_title}".

    STRICT RULES:
    1. Base your answer EXCLUSIVELY on the provided textbook pages. Do NOT invent or use external information.
    2. If the answer is NOT explicitly supported by the textbook pages, you MUST set "answer_status": "not_found".
    3. If the answer IS in the textbook:
       - Determine if the proposed answer is correct or incorrect.
       - State the verified correct answer.
       - Cite the EXACT sentence/quote from the text as "evidence_quote".
       - Specify the EXACT page number where the evidence appears.
       - Write a thorough, deep medical explanation in natural Egyptian medical student style: Arabic with proper Medical English terms for all structures/nerves/diseases.

    QUESTION:
    {question_text}

    OPTIONS (if any):
    {options_text}

    PROPOSED ANSWER (if provided):
    {proposed_answer}

    TEXTBOOK PAGES:
    {pages_context}

    Respond ONLY in strict JSON format:
    {{
        "answer_status": "correct" | "incorrect" | "not_found",
        "verified_answer": "...",
        "is_proposed_correct": true/false,
        "page_number": 42,
        "evidence_quote": "exact sentence from page text",
        "explanation_ar": "..."
    }}
    """

    client = gemini_service.get_gemini_client()
    for model in gemini_service.MODEL_PREFERENCES:
        try:
            response = client.models.generate_content(
                model=model,
                contents=prompt,
                config=gemini_service.types.GenerateContentConfig(
                    temperature=0.1,
                    response_mime_type="application/json"
                )
            )
            if response and response.text:
                if hasattr(response, 'usage_metadata') and response.usage_metadata:
                    gemini_service.log_token_usage(
                        model,
                        response.usage_metadata.prompt_token_count or 0,
                        response.usage_metadata.candidates_token_count or 0,
                        response.usage_metadata.total_token_count or 0,
                        "book_auditor_verify"
                    )
                data = json.loads(gemini_service.clean_json_text(response.text))
                status = data.get("answer_status", "not_found")

                if status == "not_found":
                    conn.close()
                    return {
                        "status": "not_found",
                        "message": "لم أجد تأكيداً أو دليلاً قاطعاً لهذه المعلومة داخل صفحات الكتاب المرفوع.",
                        "can_search_web": True
                    }

                page_num = int(data.get("page_number") or candidate_pages[0][1])
                evidence = data.get("evidence_quote", "")

                # Generate highlighted snippet crop
                snippet_res = book_snippet_generator.generate_highlighted_snippet(
                    pdf_path=pdf_path,
                    page_number=page_num,
                    evidence_quote=evidence,
                    book_id=book_id
                )

                conn.close()
                return {
                    "status": status, # 'correct' or 'incorrect'
                    "verified_answer": data.get("verified_answer", ""),
                    "is_proposed_correct": data.get("is_proposed_correct", (status == "correct")),
                    "page_number": page_num,
                    "evidence_quote": evidence,
                    "explanation": data.get("explanation_ar", ""),
                    "snippet_image_url": snippet_res.get("image_url", ""),
                    "snippet_local_path": snippet_res.get("local_path", ""),
                    "cached": False
                }
        except Exception as e:
            continue

    conn.close()
    return {"status": "error", "message": "حدث خطأ أثناء معالجة السؤال بواسطة الذكاء الاصطناعي"}

def save_verified_qa_cache(book_id: int, question_text: str, answer_text: str, explanation: str, page_number: int, snippet_image_path: str):
    """Saves confirmed convinced Q&A to cache for instant future retrieval."""
    q_hash = normalize_text_hash(question_text)
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO book_verified_qa (book_id, question_text, question_hash, answer_text, explanation, page_number, snippet_image_path, is_convinced, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, 1, datetime('now', 'localtime'))
    """, (book_id, question_text, q_hash, answer_text, explanation, page_number, snippet_image_path))
    conn.commit()
    conn.close()

# =========================================================================
# BATCH FILE AUDIT & WORD DOCUMENT PROCESSING
# =========================================================================

def parse_docx_questions(docx_path: str) -> list:
    """
    Parses a Word document containing multiple MCQs.
    Recognizes answers formatted as:
    1. 'Answer: A' or 'Ans: B' in a paragraph beneath the question.
    2. Star '*' next to the correct choice (e.g. '*A) ...' or 'B) ... *').
    Returns a list of question dicts.
    """
    doc = Document(docx_path)
    questions = []
    current_q = None

    q_start_regex = re.compile(r'^(?:Q(?:uestion)?\s*[\d\.]*|(\d+)[\.\)\-\:])\s*(.*)', re.IGNORECASE)
    option_regex = re.compile(r'^\s*([*]?)\s*([A-Ea-e])\s*[\)\.\:\-]\s*(.*?)\s*([*]?)$')
    answer_regex = re.compile(r'^(?:Answer|Ans|Correct\s*Answer)\s*[\:\-\=]\s*([A-Ea-e])', re.IGNORECASE)

    for p_idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue

        # Check for answer line
        ans_match = answer_regex.match(text)
        if ans_match:
            if current_q:
                current_q["original_answer"] = ans_match.group(1).upper()
                current_q["answer_format"] = "answer_label"
                current_q["answer_para_idx"] = p_idx
            continue

        # Check for option line
        opt_match = option_regex.match(text)
        if opt_match:
            star_before = opt_match.group(1) == '*'
            opt_letter = opt_match.group(2).upper()
            opt_content = opt_match.group(3).strip()
            star_after = opt_match.group(4) == '*'
            is_starred = star_before or star_after

            if current_q:
                current_q["options"].append({
                    "letter": opt_letter,
                    "text": opt_content,
                    "para_idx": p_idx,
                    "is_starred": is_starred,
                    "raw_text": text
                })
                if is_starred and not current_q["original_answer"]:
                    current_q["original_answer"] = opt_letter
                    current_q["answer_format"] = "star_option"
            continue

        # Check for new question stem
        q_match = q_start_regex.match(text)
        if q_match:
            # Save previous question
            if current_q and current_q["question_text"]:
                questions.append(current_q)

            current_q = {
                "question_number": len(questions) + 1,
                "question_text": text,
                "stem_para_idx": p_idx,
                "options": [],
                "original_answer": "",
                "answer_format": "unknown",
                "answer_para_idx": None
            }
        else:
            # If current_q exists and no options yet, append to stem
            if current_q and len(current_q["options"]) == 0:
                current_q["question_text"] += " " + text

    if current_q and current_q["question_text"]:
        questions.append(current_q)

    return questions

def create_audit_session(book_id: int, source_filename: str, source_file_path: str, questions: list) -> int:
    """Creates a new session in book_audit_sessions and records all parsed questions."""
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO book_audit_sessions (book_id, source_filename, source_file_path, total_questions, audited_questions, status)
        VALUES (?, ?, ?, ?, 0, 'in_progress')
    """, (book_id, source_filename, source_file_path, len(questions)))
    session_id = cur.lastrowid

    for q in questions:
        cur.execute("""
            INSERT INTO book_audit_questions (
                session_id, question_number, question_text, options_json, original_answer, status
            ) VALUES (?, ?, ?, ?, ?, 'pending')
        """, (
            session_id,
            q["question_number"],
            q["question_text"],
            json.dumps(q["options"], ensure_ascii=False),
            q["original_answer"]
        ))

    conn.commit()
    conn.close()
    return session_id

def process_next_audit_question(session_id: int) -> dict:
    """
    Audits the next pending question in the batch.
    If match -> marks verified_correct and proceeds.
    If disagreement or not_found -> pauses session and returns consultation payload.
    """
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM book_audit_sessions WHERE id=?", (session_id,))
    session = cur.fetchone()
    if not session:
        conn.close()
        return {"status": "error", "message": "الجلسة غير موجودة"}

    book_id = session["book_id"]

    # Get next pending question
    cur.execute("""
        SELECT * FROM book_audit_questions 
        WHERE session_id=? AND status='pending' 
        ORDER BY question_number ASC LIMIT 1
    """, (session_id,))
    q_row = cur.fetchone()

    if not q_row:
        # All questions processed! Finalize Word docs
        cur.execute("UPDATE book_audit_sessions SET status='completed' WHERE id=?", (session_id,))
        conn.commit()
        conn.close()
        # Generate both files
        file_paths = generate_audited_word_documents(session_id)
        return {
            "status": "completed",
            "message": "اكتمل تدقيق جميع أسئلة الملف بنجاح!",
            "clean_word_url": f"/api/book_auditor/download/{os.path.basename(file_paths['clean_word'])}",
            "report_word_url": f"/api/book_auditor/download/{os.path.basename(file_paths['report_word'])}"
        }

    q_id = q_row["id"]
    q_num = q_row["question_number"]
    q_text = q_row["question_text"]
    orig_ans = q_row["original_answer"]
    options = json.loads(q_row["options_json"] or "[]")
    options_formatted = "\n".join([f"{o.get('letter', '')}) {o.get('text', '')}" for o in options])

    conn.close()

    # Verify against book
    verify_res = verify_single_question_against_book(
        book_id=book_id,
        question_text=q_text,
        options_text=options_formatted,
        proposed_answer=orig_ans
    )

    conn = database.get_connection()
    cur = conn.cursor()

    if verify_res.get("status") == "not_found":
        # Pause session and consult student
        cur.execute("""
            UPDATE book_audit_sessions SET status='paused_waiting_user' WHERE id=?
        """, (session_id,))
        cur.execute("""
            UPDATE book_audit_questions SET
                status='not_found',
                explanation=?,
                consultation_chat_json=?
            WHERE id=?
        """, (
            verify_res.get("message", "لم توجد في الكتاب"),
            json.dumps([{"sender": "ai", "text": "لم أجد تأكيداً لهذه المعلومة داخل صفحات الكتاب. هل تفضل أن أعيد البحث بمصطلحات أخرى في الكتاب أم أبحث عبر الإنترنت؟"}], ensure_ascii=False),
            q_id
        ))
        conn.commit()
        conn.close()

        return {
            "status": "paused_not_found",
            "question_id": q_id,
            "question_number": q_num,
            "question_text": q_text,
            "options": options,
            "original_answer": orig_ans,
            "message": "لم أتمكن من إيجاد المعلومة في الكتاب المرفوع."
        }

    elif verify_res.get("status") == "incorrect" or (not verify_res.get("is_proposed_correct", True)):
        # Found incorrect in the book! Pause session and consult student
        proposed_correct = verify_res.get("verified_answer", "")
        cur.execute("""
            UPDATE book_audit_sessions SET status='paused_waiting_user' WHERE id=?
        """, (session_id,))
        cur.execute("""
            UPDATE book_audit_questions SET
                status='found_incorrect',
                proposed_answer=?,
                explanation=?,
                page_number=?,
                snippet_image_path=?,
                consultation_chat_json=?
            WHERE id=?
        """, (
            proposed_correct,
            verify_res.get("explanation", ""),
            verify_res.get("page_number", 0),
            verify_res.get("snippet_image_url", ""),
            json.dumps([{"sender": "ai", "text": f"الإجابة المحددة في الملف هي ({orig_ans}) ولكن الكتاب في صفحة ({verify_res.get('page_number')}) ينص على أن الإجابة الصحيحة هي ({proposed_correct}) لأن: {verify_res.get('explanation')} هل تريد تعديل الإجابة في ملف الوورد إلى ({proposed_correct}) أم الإبقاء عليها؟"}], ensure_ascii=False),
            q_id
        ))
        conn.commit()
        conn.close()

        return {
            "status": "paused_disagreement",
            "question_id": q_id,
            "question_number": q_num,
            "question_text": q_text,
            "options": options,
            "original_answer": orig_ans,
            "proposed_answer": proposed_correct,
            "page_number": verify_res.get("page_number", 0),
            "snippet_image_url": verify_res.get("snippet_image_url", ""),
            "explanation": verify_res.get("explanation", "")
        }

    else:
        # Verified correct!
        cur.execute("""
            UPDATE book_audit_questions SET
                status='verified_correct',
                final_answer=?,
                explanation=?,
                page_number=?,
                snippet_image_path=?
            WHERE id=?
        """, (
            orig_ans,
            verify_res.get("explanation", ""),
            verify_res.get("page_number", 0),
            verify_res.get("snippet_image_url", ""),
            q_id
        ))
        cur.execute("""
            UPDATE book_audit_sessions SET audited_questions = audited_questions + 1 WHERE id=?
        """, (session_id,))
        conn.commit()
        conn.close()

        return {
            "status": "question_verified_correct",
            "question_number": q_num,
            "audited_questions": session["audited_questions"] + 1,
            "total_questions": session["total_questions"]
        }

def resolve_consultation_question(session_id: int, question_id: int, user_action: str, target_answer: str = "") -> dict:
    """
    Resolves a paused question after consulting with the user:
    - 'modify': accept the AI's proposed answer.
    - 'keep': keep the original file's answer.
    - 'web_search': execute internet search for this question and update.
    """
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM book_audit_questions WHERE id=? AND session_id=?", (question_id, session_id))
    q = cur.fetchone()
    if not q:
        conn.close()
        return {"status": "error", "message": "السؤال غير موجود"}

    cur.execute("SELECT * FROM book_audit_sessions WHERE id=?", (session_id,))
    session = cur.fetchone()
    book_id = session["book_id"]

    if user_action == "web_search":
        conn.close()
        # Run web search for this question
        res = verify_single_question_against_book(
            book_id=book_id,
            question_text=q["question_text"],
            options_text=q["options_json"],
            proposed_answer=q["original_answer"],
            force_web_search=True
        )
        conn = database.get_connection()
        cur = conn.cursor()
        cur.execute("""
            UPDATE book_audit_questions SET
                status='web_consulted',
                proposed_answer=?,
                explanation=?
            WHERE id=?
        """, (
            res.get("verified_answer", q["original_answer"]),
            res.get("explanation", "") + "\n\nالمصدر: " + res.get("source_reference", ""),
            question_id
        ))
        conn.commit()
        conn.close()
        return {
            "status": "web_search_complete",
            "proposed_answer": res.get("verified_answer", q["original_answer"]),
            "explanation": res.get("explanation", ""),
            "source": res.get("source_reference", "")
        }

    final_ans = target_answer if target_answer else (q["proposed_answer"] if user_action == "modify" else q["original_answer"])
    new_status = "confirmed_modified" if user_action == "modify" else "confirmed_kept"

    cur.execute("""
        UPDATE book_audit_questions SET
            status=?,
            final_answer=?
        WHERE id=?
    """, (new_status, final_ans, question_id))

    cur.execute("""
        UPDATE book_audit_sessions SET
            audited_questions = audited_questions + 1,
            status = 'in_progress'
        WHERE id=?
    """, (session_id,))
    conn.commit()
    conn.close()

    return {"status": "resumed", "final_answer": final_ans}

def generate_audited_word_documents(session_id: int) -> dict:
    """
    Creates TWO Word documents:
    1. Clean Modified Word: Same format as user's docx, ONLY answers modified where agreed.
    2. Comprehensive Evidence Report: Contains questions, explanations, and embedded highlighted snippet screenshots.
    """
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("SELECT * FROM book_audit_sessions WHERE id=?", (session_id,))
    session = cur.fetchone()
    cur.execute("SELECT * FROM reference_books WHERE id=?", (session["book_id"],))
    book = cur.fetchone()
    cur.execute("SELECT * FROM book_audit_questions WHERE session_id=? ORDER BY question_number ASC", (session_id,))
    questions = cur.fetchall()
    conn.close()

    orig_path = session["source_file_path"]
    book_title = book["title"]
    base_name = os.path.splitext(session["source_filename"])[0]

    # 1. GENERATE CLEAN MODIFIED DOCX
    clean_docx_filename = f"Audited_Clean_{base_name}.docx"
    clean_docx_path = os.path.join(AUDIT_DOCS_DIR, clean_docx_filename)

    if orig_path and os.path.exists(orig_path):
        clean_doc = Document(orig_path)
    else:
        clean_doc = Document()

    # Map question modifications
    modified_qs = {}
    for q in questions:
        orig = (q["original_answer"] or "").strip().upper()
        final = (q["final_answer"] or orig).strip().upper()
        if final and final != orig and q["status"] == "confirmed_modified":
            modified_qs[q["question_number"]] = {
                "old": orig,
                "new": final
            }

    # If we opened original doc, update only modified answer paragraphs/stars
    if orig_path and os.path.exists(orig_path) and modified_qs:
        ans_pattern = re.compile(r'^(Answer|Ans|Correct\s*Answer)\s*[\:\-\=]\s*([A-Ea-e])', re.IGNORECASE)
        opt_pattern = re.compile(r'^(\s*)([*]?\s*)([A-Ea-e])(\s*[\)\.\:\-]\s*)(.*?)(\s*[*]?)$')

        current_q_num = 0
        for p in clean_doc.paragraphs:
            text = p.text.strip()
            # Check question start
            q_match = re.match(r'^(?:Q(?:uestion)?\s*[\d\.]*|(\d+)[\.\)\-\:])', text, re.IGNORECASE)
            if q_match:
                current_q_num += 1

            if current_q_num in modified_qs:
                mod = modified_qs[current_q_num]
                # Check for Answer: A
                ans_m = ans_pattern.match(text)
                if ans_m:
                    p.text = f"{ans_m.group(1)}: {mod['new']}"
                # Check for *A)
                opt_m = opt_pattern.match(p.text)
                if opt_m:
                    letter = opt_m.group(3).upper()
                    content = opt_m.group(5)
                    # If this was old answer, remove star
                    if letter == mod["old"]:
                        p.text = f"{letter}) {content}"
                    # If this is new answer, add star
                    elif letter == mod["new"]:
                        p.text = f"*{letter}) {content}"

    clean_doc.save(clean_docx_path)

    # 2. GENERATE COMPREHENSIVE EVIDENCE & HIGHLIGHT REPORT
    report_docx_filename = f"Evidence_Report_{base_name}.docx"
    report_docx_path = os.path.join(AUDIT_DOCS_DIR, report_docx_filename)

    rep_doc = Document()
    for sec in rep_doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # Title
    t_para = rep_doc.add_heading(level=0)
    t_run = t_para.add_run(f"📖 تقرير التدقيق المرجعي الشامل والأدلة الموثقة")
    t_run.font.color.rgb = RGBColor(13, 148, 136)

    # Meta
    p_meta = rep_doc.add_paragraph()
    p_meta.add_run(f"الكتاب المرجعي: {book_title}\n").bold = True
    p_meta.add_run(f"الملف الأصلي: {session['source_filename']}\n")
    p_meta.add_run(f"تاريخ التدقيق: {time.strftime('%Y-%m-%d %H:%M')}\n")
    p_meta.add_run(f"إجمالي الأسئلة: {len(questions)} سؤال | الأسئلة المعدلة: {len(modified_qs)}\n")

    rep_doc.add_paragraph("―" * 50)

    for q in questions:
        q_num = q["question_number"]
        q_text = q["question_text"]
        orig_ans = q["original_answer"]
        final_ans = q["final_answer"] or orig_ans
        status = q["status"]
        explanation = q["explanation"] or ""
        page_num = q["page_number"]
        snippet_rel = q["snippet_image_path"]

        # Question Title
        h = rep_doc.add_heading(level=2)
        h.add_run(f"السؤال {q_num}: {q_text}")

        # Status & Answers Box
        p_ans = rep_doc.add_paragraph()
        if status == "confirmed_modified":
            p_ans.add_run(f"⚠️ الإجابة الأصلية في الملف: ({orig_ans}) ➔ تم التعديل إلى: ({final_ans}) ✅\n").bold = True
            p_ans.runs[0].font.color.rgb = RGBColor(217, 119, 6)
        else:
            p_ans.add_run(f"✅ الإجابة المعتمدة: ({final_ans})\n").bold = True
            p_ans.runs[0].font.color.rgb = RGBColor(16, 185, 129)

        if page_num > 0:
            p_ans.add_run(f"📄 موضع الدليل في الكتاب: صفحة {page_num}\n").bold = True

        # Explanation
        if explanation:
            p_exp = rep_doc.add_paragraph()
            p_exp.add_run("💡 التفسير والتعليل الطبي المعتمد من الكتاب:\n").bold = True
            p_exp.add_run(explanation)

        # Embed Snippet Image if available
        if snippet_rel:
            img_full_path = os.path.join(BASE_DIR, snippet_rel.lstrip("/"))
            if os.path.exists(img_full_path):
                try:
                    rep_doc.add_paragraph("🔍 لقطة الشاشة مع التظليل (Highlight) من صفحة المرجع:")
                    rep_doc.add_picture(img_full_path, width=Inches(5.6))
                except Exception as img_err:
                    pass

        rep_doc.add_paragraph("―" * 40)

    rep_doc.save(report_docx_path)

    # Save paths to DB
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("""
        UPDATE book_audit_sessions SET
            clean_word_path=?,
            report_word_path=?
        WHERE id=?
    """, (clean_docx_path, report_docx_path, session_id))
    conn.commit()
    conn.close()

    return {
        "clean_word": clean_docx_path,
        "report_word": report_docx_path
    }
