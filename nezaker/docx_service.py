import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
WORD_BANKS_DIR = os.path.join(BASE_DIR, "lecture_word_banks")
os.makedirs(WORD_BANKS_DIR, exist_ok=True)

def sanitize_filename(name: str) -> str:
    cleaned = re.sub(r'[\\/*?:"<>|]', '', name).strip()
    return cleaned[:80]

def get_word_file_path(lecture_id: int, lecture_number: int, lecture_title: str) -> str:
    num_prefix = f"{lecture_number:02d}_" if lecture_number > 0 else f"lec_{lecture_id}_"
    clean_title = sanitize_filename(lecture_title)
    filename = f"{num_prefix}{clean_title}_Questions.docx"
    return os.path.join(WORD_BANKS_DIR, filename)

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def clean_text_content(val: str) -> str:
    if not val:
        return ""
    text = str(val).strip()
    return re.sub(r'^[•\-\*\d\.\)\(\s]+', '', text).strip()

def clean_option_content(val: str, letter: str) -> str:
    if not val:
        return ""
    text = str(val).strip()
    # Remove leading option labels like "a.", "A)", "(A)", "A -", etc.
    text = re.sub(rf'^\(?[{letter}{letter.lower()}][\.\)\:\-\s]+\s*', '', text, flags=re.IGNORECASE).strip()
    return text

def rebuild_lecture_word_document(lecture_id: int, selected_batches: list = None, output_path: str = None, include_explanation: bool = True) -> str:
    """
    Completely rebuilds or exports the lecture's Word document from SQLite questions.
    If selected_batches is provided (list of batch numbers), only questions from those batches are included.
    Formats questions in clean format:
    1. question
    a. answer
    b. answer
    c. answer
    d. answer
    ? explanation (if include_explanation is True)
    Returns the saved file path.
    """
    import database
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("SELECT id, lecture_number, title, subject FROM lectures WHERE id=?", (lecture_id,))
    lec = cur.fetchone()
    if not lec:
        conn.close()
        raise ValueError(f"Lecture {lecture_id} not found")

    lec_number = lec["lecture_number"]
    lec_title = lec["title"]

    if selected_batches is not None and len(selected_batches) > 0:
        placeholders = ",".join("?" * len(selected_batches))
        cur.execute(f"""
            SELECT * FROM questions 
            WHERE lecture_id=? AND batch_number IN ({placeholders})
            ORDER BY batch_number ASC, id ASC
        """, [lecture_id] + list(selected_batches))
    else:
        cur.execute("""
            SELECT * FROM questions 
            WHERE lecture_id=? 
            ORDER BY batch_number ASC, id ASC
        """, (lecture_id,))

    raw_questions = [dict(r) for r in cur.fetchall()]
    conn.close()

    # Determine file path
    if not output_path:
        file_path = get_word_file_path(lecture_id, lec_number, lec_title)
    else:
        file_path = output_path

    doc = Document()
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # Title Header
    title = doc.add_heading(level=0)
    run = title.add_run(f"📚 بنك أسئلة: {lec_title}")
    run.font.color.rgb = RGBColor(13, 148, 136) # Teal
    run.font.size = Pt(22)
    run.font.name = 'Calibri'
    
    p = doc.add_paragraph()
    p_run = p.add_run(f"منصة نُذاكِر (Nezaker) الطبية الذكية — بنك الأسئلة والحالات السريرية المجمعة ({len(raw_questions)} سؤال)")
    p_run.font.color.rgb = RGBColor(100, 116, 139)
    p_run.font.size = Pt(11)
    doc.add_paragraph("=" * 60)

    # Group by batch_number
    from collections import OrderedDict
    batches = OrderedDict()
    for q in raw_questions:
        b_num = q.get("batch_number") or 1
        b_name = q.get("batch_name") or f"باتش {b_num}"
        key = (b_num, b_name)
        if key not in batches:
            batches[key] = []
        batches[key].append(q)

    global_idx = 1
    for (b_num, b_name), b_questions in batches.items():
        batch_heading = doc.add_heading(level=2)
        b_run = batch_heading.add_run(f"⚡ دفعة #{b_num}: {b_name} ({len(b_questions)} سؤال)")
        b_run.font.color.rgb = RGBColor(37, 99, 235) # Sapphire Blue
        b_run.font.size = Pt(15)

        for q in b_questions:
            is_case = q.get("question_type") == "case" or (q.get("case_scenario") and len(q.get("case_scenario")) > 5)
            
            # Case scenario box (if any)
            if is_case and q.get("case_scenario"):
                table = doc.add_table(rows=1, cols=1)
                table.autofit = False
                table.columns[0].width = Inches(6.5)
                cell = table.cell(0, 0)
                set_cell_background(cell, "F1F5F9")
                cp = cell.paragraphs[0]
                cp_run = cp.add_run(q.get("case_scenario").strip())
                cp_run.font.size = Pt(10.5)
                cp_run.font.italic = True
                doc.add_paragraph()

            # 1. Question Prompt
            q_text = clean_text_content(q.get("question_text", ""))
            q_p = doc.add_paragraph()
            q_run = q_p.add_run(f"{global_idx}. {q_text}")
            q_run.bold = True
            q_run.font.size = Pt(11.5)
            q_run.font.color.rgb = RGBColor(15, 23, 42)

            # Options
            correct_letter = str(q.get("correct_option") or q.get("correct_answer") or "").strip().upper()
            if len(correct_letter) > 1:
                m = re.search(r'([A-D])', correct_letter)
                if m:
                    correct_letter = m.group(1)

            opts = [
                ("a", clean_option_content(q.get("option_a", ""), "a"), "A"),
                ("b", clean_option_content(q.get("option_b", ""), "b"), "B"),
                ("c", clean_option_content(q.get("option_c", ""), "c"), "C"),
                ("d", clean_option_content(q.get("option_d", ""), "d"), "D")
            ]
            for letter_lower, text, letter_upper in opts:
                if not text:
                    continue
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Inches(0.25)
                is_correct = (letter_upper == correct_letter)

                if is_correct:
                    # Highlight or star correct answer
                    opt_run = opt_p.add_run(f"{letter_lower}. {text}  *")
                    opt_run.bold = True
                    opt_run.font.color.rgb = RGBColor(22, 163, 74)
                else:
                    opt_run = opt_p.add_run(f"{letter_lower}. {text}")
                    opt_run.font.color.rgb = RGBColor(30, 41, 59)

            # Explanation if requested
            if include_explanation:
                exp_text = (q.get("explanation_arabic") or q.get("explanation") or "").strip()
                lec_ev = (q.get("lecture_evidence") or "").strip()
                full_exp = exp_text
                if lec_ev:
                    full_exp = f"{full_exp} (موضع المعلومة بالمنهج: {lec_ev})" if full_exp else f"موضع المعلومة بالمنهج: {lec_ev}"
                
                exp_p = doc.add_paragraph()
                exp_p.paragraph_format.left_indent = Inches(0.25)
                exp_run = exp_p.add_run(f"? {full_exp}" if full_exp else "?")
                exp_run.font.italic = True
                exp_run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
                exp_run.font.size = Pt(10.5)

            doc.add_paragraph() # Spacing between questions
            global_idx += 1

    try:
        doc.save(file_path)
        return file_path
    except PermissionError:
        from datetime import datetime
        base, ext = os.path.splitext(file_path)
        t_stamp = datetime.now().strftime("%H%M%S")
        alt_path = f"{base}_new_{t_stamp}{ext}"
        doc.save(alt_path)
        return alt_path

def append_questions_to_word(lecture_id: int, lecture_number: int = 0, lecture_title: str = "", questions: list = None, batch_name: str = "AI Generation", batch_number: int = None, include_explanation: bool = True) -> str:
    """Synchronously ensures the lecture's dedicated Word document contains all questions."""
    return rebuild_lecture_word_document(lecture_id, include_explanation=include_explanation)

def export_filtered_exam_questions_docx(questions: list, title: str = "أسئلة امتحانات سابقة مفلترة بالذكاء الاصطناعي", include_explanation: bool = True) -> str:
    """Creates a standalone Word document for filtered exam questions in 1. question / a. answer format."""
    from datetime import datetime
    t_stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"Exam_Filtered_{t_stamp}.docx"
    file_path = os.path.join(WORD_BANKS_DIR, filename)
    
    doc = Document()
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # Title Header
    title_heading = doc.add_heading(level=0)
    run = title_heading.add_run(f"🩺 {title}")
    run.font.color.rgb = RGBColor(13, 148, 136) # Teal
    run.font.size = Pt(22)
    run.font.name = 'Calibri'
    
    p = doc.add_paragraph()
    p_run = p.add_run(f"منصة نُذاكِر (Nezaker) الطبية الذكية — استخراج وتصفية أسئلة الامتحانات بالـ AI والـ OCR ({len(questions)} سؤال)")
    p_run.font.color.rgb = RGBColor(100, 116, 139)
    p_run.font.size = Pt(11)
    doc.add_paragraph("=" * 60)
    
    for i, q in enumerate(questions, 1):
        q_text = clean_text_content(q.get("question_text", ""))
        q_p = doc.add_paragraph()
        q_run = q_p.add_run(f"{i}. {q_text}")
        q_run.bold = True
        q_run.font.size = Pt(11.5)
        q_run.font.color.rgb = RGBColor(15, 23, 42)
        
        correct_letter = str(q.get("correct_option") or q.get("correct_answer") or "").strip().upper()
        if len(correct_letter) > 1:
            m = re.search(r'([A-D])', correct_letter)
            if m:
                correct_letter = m.group(1)
                
        opts = [
            ("a", clean_option_content(q.get("option_a", ""), "a"), "A"),
            ("b", clean_option_content(q.get("option_b", ""), "b"), "B"),
            ("c", clean_option_content(q.get("option_c", ""), "c"), "C"),
            ("d", clean_option_content(q.get("option_d", ""), "d"), "D")
        ]
        for letter_lower, text, letter_upper in opts:
            if not text:
                continue
            opt_p = doc.add_paragraph()
            opt_p.paragraph_format.left_indent = Inches(0.25)
            is_correct = (letter_upper == correct_letter)
            if is_correct:
                star_run = opt_p.add_run(f"{letter_lower}. {text}  *")
                star_run.bold = True
                star_run.font.color.rgb = RGBColor(22, 163, 74)
            else:
                opt_run = opt_p.add_run(f"{letter_lower}. {text}")
                opt_run.font.color.rgb = RGBColor(30, 41, 59)
                
        # Explanation if requested
        if include_explanation:
            exp_text = (q.get("explanation_arabic") or q.get("explanation") or "").strip()
            lec_ev = (q.get("lecture_evidence") or "").strip()
            full_exp = exp_text
            if lec_ev:
                full_exp = f"{full_exp} (موضع المعلومة بالمنهج: {lec_ev})" if full_exp else f"موضع المعلومة بالمنهج: {lec_ev}"
            
            exp_p = doc.add_paragraph()
            exp_p.paragraph_format.left_indent = Inches(0.25)
            exp_run = exp_p.add_run(f"? {full_exp}" if full_exp else "?")
            exp_run.font.italic = True
            exp_run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
            exp_run.font.size = Pt(10.5)
            
        doc.add_paragraph() # Spacing between questions
        
    try:
        doc.save(file_path)
        return file_path
    except PermissionError:
        base, ext = os.path.splitext(file_path)
        alt_path = f"{base}_new_{datetime.now().strftime('%H%M%S')}{ext}"
        doc.save(alt_path)
        return alt_path

def export_comprehensive_mock_exam_docx(title: str, mcq_questions: list, essay_questions: list, lecture_title: str = "") -> str:
    """
    Generates a beautifully formatted Microsoft Word document for a complete mastery mock exam,
    including MCQs, Clinical Cases, and Short-Answer / Essay Questions with Model Answers.
    """
    exam_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scratch_uploads", "mock_exams")
    os.makedirs(exam_dir, exist_ok=True)
    
    clean_title = re.sub(r'[\\/*?:"<>|]', "", title).strip().replace(" ", "_")[:60]
    file_path = os.path.join(exam_dir, f"{clean_title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")

    doc = Document()
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # Header
    head_p = doc.add_paragraph()
    head_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = head_p.add_run(f"🏥 {title}\n")
    hr.bold = True
    hr.font.size = Pt(18)
    hr.font.color.rgb = RGBColor(30, 58, 138)
    
    if lecture_title:
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sub_p.add_run(f"المحاضرة: {lecture_title} | تاريخ الامتحان: {datetime.now().strftime('%Y-%m-%d')}")
        sr.font.size = Pt(11)
        sr.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("-" * 50)

    # ----------------- PART I: MCQs & CASES -----------------
    if mcq_questions and len(mcq_questions) > 0:
        p1 = doc.add_paragraph()
        p1_run = p1.add_run(f"📌 القسم الأول: أسئلة الاختيار من متعدد والحالات السريرية ({len(mcq_questions)} سؤال)")
        p1_run.bold = True
        p1_run.font.size = Pt(14)
        p1_run.font.color.rgb = RGBColor(15, 23, 42)

        for idx, q in enumerate(mcq_questions, 1):
            qp = doc.add_paragraph()
            q_run = qp.add_run(f"سؤال [{idx}]: ")
            q_run.bold = True
            q_run.font.color.rgb = RGBColor(37, 99, 235)

            if q.get("case_scenario"):
                sc_run = qp.add_run(f"[حالة سريرية] {q.get('case_scenario')}\n")
                sc_run.italic = True

            qp.add_run(q.get("question_text", ""))

            correct_letter = (q.get("correct_option") or "A").strip().upper()
            opts = [
                ("A", q.get("option_a", "")),
                ("B", q.get("option_b", "")),
                ("C", q.get("option_c", "")),
                ("D", q.get("option_d", ""))
            ]
            for letter, text in opts:
                if not text:
                    continue
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Inches(0.25)
                is_correct = (letter == correct_letter)
                if is_correct:
                    opt_p.add_run(f"✔ ({letter}) {text}  *").bold = True
                else:
                    opt_p.add_run(f"   ({letter}) {text}")

            # Explanation
            exp_text = q.get("explanation_arabic") or q.get("explanation") or ""
            if exp_text:
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F0FDF4")
                ep = cell.paragraphs[0]
                ep.add_run(f"✔ الإجابة الصحيحة: ({correct_letter})\n").bold = True
                ep.add_run(f"التعليل الطبي: {exp_text}")

            doc.add_paragraph()

    # ----------------- PART II: ESSAY QUESTIONS -----------------
    if essay_questions and len(essay_questions) > 0:
        doc.add_page_break()
        p2 = doc.add_paragraph()
        p2_run = p2.add_run(f"📝 القسم الثاني: الأسئلة المقالية والتحليل السريري ({len(essay_questions)} أسئلة)")
        p2_run.bold = True
        p2_run.font.size = Pt(14)
        p2_run.font.color.rgb = RGBColor(126, 34, 206)

        inst = doc.add_paragraph()
        inst.add_run("تعليمات: أجب عن الأسئلة المقالية التالية في المساحات المخصصة، ثم طابق إجابتك مع معايير التصحيح والإجابة النموذجية بالأسفل.").italic = True

        for idx, eq in enumerate(essay_questions, 1):
            ep = doc.add_paragraph()
            eq_run = ep.add_run(f"السؤال المقالي [{idx}]: ")
            eq_run.bold = True
            eq_run.font.color.rgb = RGBColor(126, 34, 206)

            if eq.get("clinical_context"):
                ep.add_run(f"({eq.get('clinical_context')})\n").italic = True

            ep.add_run(eq.get("question_text", ""))
            
            # Blank lines for student to write
            doc.add_paragraph("\n".join(["_" * 65] * 3))

            # Model Answer & Rubric
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "FAF5FF")
            cp = cell.paragraphs[0]
            cp.add_run("🎯 الإجابة النموذجية ومعايير التصحيح (Model Answer & Rubric):\n").bold = True
            cp.add_run(eq.get("ideal_model_answer", "") + "\n\n")

            key_pts = eq.get("key_points_to_mention", [])
            if key_pts:
                cp.add_run("النقاط الجوهرية الواجب ذكرها لنيل الدرجة الكاملة:\n").bold = True
                for pt in key_pts:
                    cp.add_run(f" • {pt}\n")

            doc.add_paragraph("-" * 40)

    try:
        doc.save(file_path)
        return file_path
    except PermissionError:
        base, ext = os.path.splitext(file_path)
        alt_path = f"{base}_alt_{datetime.now().strftime('%H%M%S')}{ext}"
        doc.save(alt_path)
        return alt_path

def open_word_document(file_path: str):
    """Launches the Word document in Microsoft Word on Windows."""
    try:
        if os.path.exists(file_path):
            os.startfile(file_path)
            return True
        return False
    except Exception as e:
        print(f"Error opening Word doc: {e}")
        return False

if __name__ == "__main__":
    print("DOCX Service ready. Word banks directory:", WORD_BANKS_DIR)
